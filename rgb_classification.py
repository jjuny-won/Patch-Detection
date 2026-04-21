"""
RGB Color Classification — pH & Temperature
RandomForest + Feature Engineering
------------------------------------------------------------
[pH Baseline]  Accuracy: 95.70%  (pH5.10 F1: 0.921, pH9 F1: 0.891)
[pH Enhanced]  Accuracy: 95.92%  (pH5.10 F1: 0.924, pH9 F1: 0.895)
[Temp]         Accuracy: 98.12%  (모든 구간 우수)
"""

import numpy as np
import pandas as pd
import openpyxl
import warnings
from sklearn.ensemble import RandomForestClassifier
from sklearn.model_selection import train_test_split
from sklearn.metrics import accuracy_score, classification_report, confusion_matrix

warnings.filterwarnings("ignore")

# ──────────────────────────────────────────────────────────────
# 1. 데이터 로드
# ──────────────────────────────────────────────────────────────

def load_ph_data(filepath: str) -> pd.DataFrame:
    """
    pH 엑셀 파일 로드.
    컬럼 순서: pH5(0), pH6(3), pH7(6), pH8(9), pH9(12), pH510(15)
    """
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active

    label_col_map = {
        "pH5": 0, "pH6": 3, "pH7": 6,
        "pH8": 9, "pH9": 12, "pH510": 15,
    }
    rows = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        for label, col in label_col_map.items():
            if row[col] is not None:
                try:
                    r, g, b = float(row[col]), float(row[col+1]), float(row[col+2])
                    rows.append({"R": r, "G": g, "B": b, "label": label})
                except (TypeError, ValueError):
                    pass
    return pd.DataFrame(rows)


def load_temp_data(filepath: str) -> pd.DataFrame:
    """
    온도 엑셀 파일 로드.
    컬럼 순서: 33°C(0), 34°C(3), 35°C(6), 36°C(9), 37°C(12),
               38°C(15), 39°C(18), 40°C(21)
    """
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active

    label_col_map = {
        "33C": 0,  "34C": 3,  "35C": 6,  "36C": 9,
        "37C": 12, "38C": 15, "39C": 18, "40C": 21,
    }
    rows = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0] is None:
            continue
        for label, col in label_col_map.items():
            if len(row) > col + 2 and row[col] is not None:
                try:
                    r, g, b = float(row[col]), float(row[col+1]), float(row[col+2])
                    rows.append({"R": r, "G": g, "B": b, "label": label})
                except (TypeError, ValueError):
                    pass
    return pd.DataFrame(rows)


# ──────────────────────────────────────────────────────────────
# 2. Feature Engineering
# ──────────────────────────────────────────────────────────────

def add_features(df: pd.DataFrame) -> pd.DataFrame:
    """
    RGB 3채널 → 19개 feature로 확장.

    주요 개선 근거 (pH5.10 vs pH9 오버랩):
      - pH5.10: R=10.8, G=22.3, B=30.4  → R-G diff ≈ -11.5
      - pH9:    R=5.7,  G=8.5,  B=16.6  → R-G diff ≈ -2.8
      → R_minus_G 가 가장 큰 판별력 (feature importance 1위, 0.188)
      → chroma, hue, G_B ratio 보조
    """
    df = df.copy()
    R, G, B = df["R"], df["G"], df["B"]
    eps = 1e-6

    # ① 채널 간 차이 — pH5.10 vs pH9 핵심 구분자
    df["R_minus_G"] = R - G                     # importance 0.188 ★
    df["R_minus_B"] = R - B                     # importance 0.073
    df["G_minus_B"] = G - B                     # importance 0.105

    # ② 채널 비율 — 밝기 불변, 색상 비율 보존
    df["R_G"] = R / (G + eps)                   # importance 0.079
    df["R_B"] = R / (B + eps)                   # importance 0.052
    df["G_B"] = G / (B + eps)

    # ③ 밝기·명도
    df["intensity"]   = (R + G + B) / 3.0
    df["max_ch"]      = np.maximum(np.maximum(R, G), B)
    df["min_ch"]      = np.minimum(np.minimum(R, G), B)
    df["chroma"]      = df["max_ch"] - df["min_ch"]  # HSV saturation proxy

    # ④ 정규화 RGB — 밝기 제거 후 순수 색상 비율
    total = R + G + B + eps
    df["r_norm"] = R / total
    df["g_norm"] = G / total
    df["b_norm"] = B / total

    # ⑤ Hue (HSV, 0~360°)
    r_n = R.values / 255.0
    g_n = G.values / 255.0
    b_n = B.values / 255.0
    maxc  = np.maximum(np.maximum(r_n, g_n), b_n)
    minc  = np.minimum(np.minimum(r_n, g_n), b_n)
    delta = maxc - minc + eps
    hue = np.where(
        maxc == r_n, (g_n - b_n) / delta % 6,
        np.where(maxc == g_n, (b_n - r_n) / delta + 2,
                              (r_n - g_n) / delta + 4)
    ) * 60 % 360
    df["hue"] = np.where(delta < eps, 0.0, hue)    # importance 0.095

    # ⑥ HSV Saturation
    df["saturation"] = df["chroma"] / (df["max_ch"] + eps)

    # ⑦ Luminance proxy (ITU-R BT.601)
    df["L_proxy"] = 0.299 * R + 0.587 * G + 0.114 * B

    return df


FEATURE_COLS = [
    "R", "G", "B",
    "R_minus_G", "R_minus_B", "G_minus_B",
    "R_G", "R_B", "G_B",
    "intensity", "max_ch", "min_ch", "chroma",
    "r_norm", "g_norm", "b_norm",
    "hue", "saturation", "L_proxy",
]


# ──────────────────────────────────────────────────────────────
# 3. 학습 & 평가
# ──────────────────────────────────────────────────────────────

def train_and_evaluate(df: pd.DataFrame, label: str) -> RandomForestClassifier:
    df = add_features(df)
    X = df[FEATURE_COLS].values
    y = df["label"].values

    X_tr, X_te, y_tr, y_te = train_test_split(
        X, y, test_size=0.2, random_state=42, stratify=y
    )

    # Baseline (RGB only)
    rf_base = RandomForestClassifier(n_estimators=100, random_state=42, n_jobs=-1)
    rf_base.fit(X_tr[:, :3], y_tr)
    acc_base = accuracy_score(y_te, rf_base.predict(X_te[:, :3]))

    # Enhanced (19 features)
    rf = RandomForestClassifier(
        n_estimators=200,
        max_depth=None,
        min_samples_leaf=2,
        random_state=42,
        n_jobs=-1,
    )
    rf.fit(X_tr, y_tr)
    y_pred = rf.predict(X_te)
    acc_enh = accuracy_score(y_te, y_pred)

    print(f"\n{'='*55}")
    print(f"  {label}")
    print(f"{'='*55}")
    print(f"  Baseline (RGB only) Accuracy : {acc_base*100:.2f}%")
    print(f"  Enhanced (19 feat)  Accuracy : {acc_enh*100:.2f}%  "
          f"(+{(acc_enh-acc_base)*100:.2f}%p)")
    print()
    print(classification_report(y_te, y_pred, digits=4))

    # Feature importance
    imp = pd.Series(rf.feature_importances_, index=FEATURE_COLS)
    imp = imp.sort_values(ascending=False)
    print("  [Feature Importance Top 10]")
    for fname, val in imp.head(10).items():
        bar = "█" * int(val * 200)
        print(f"  {fname:<15} {val:.4f}  {bar}")

    return rf


# ──────────────────────────────────────────────────────────────
# 4. 실행
# ──────────────────────────────────────────────────────────────

if __name__ == "__main__":
    # 파일 경로 수정
    PH_FILE   = "pH_색상변화.xlsx"
    TEMP_FILE = "온도_색상변화__1_.xlsx"

    print("Loading data...")
    df_ph   = load_ph_data(PH_FILE)
    df_temp = load_temp_data(TEMP_FILE)

    print(f"pH data  : {len(df_ph):,} samples, classes: {sorted(df_ph['label'].unique())}")
    print(f"Temp data: {len(df_temp):,} samples, classes: {sorted(df_temp['label'].unique())}")

    rf_ph   = train_and_evaluate(df_ph,   label="pH Classification")
    rf_temp = train_and_evaluate(df_temp, label="Temperature Classification")
